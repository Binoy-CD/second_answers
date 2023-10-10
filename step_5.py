from selenium import webdriver
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import Select
from selenium.common.exceptions import TimeoutException
from selenium.webdriver.common.action_chains import ActionChains
import gspread
import vedantu
import toppr
import byjus
import embibe
import pandas as pd
from bs4 import BeautifulSoup
import time
import bs4
import re
from bs4 import BeautifulSoup
import openpyxl
import bs4
import logging
import warnings
warnings.filterwarnings("ignore")

logging.basicConfig( filename='/home/binoy/second_answers/log_of_process.log',
                    format='%(asctime)s %(levelname)s %(threadName)-10s %(message)s',
                     level=logging.INFO)
logger = logging.getLogger(__name__)

def initiate_driver():
    logging.info('Initiating driver...')
    global driver
    PATH = "/usr/bin/chromedriver"
    py = "127.0.0.1:24001"
    # service = Service(executable_path=PATH)
    options = Options()
    options.add_argument('--proxy-server=%s' % py)
    options.add_argument('--ignore-ssl-errors=yes')
    options.add_argument('--ignore-certificate-errors')
    options.add_argument('--allow-insecure-localhost')
    options.add_argument('headless')
    options.add_argument('--no-sandbox') 
    options.add_argument('--disable-gpu')
    options.add_argument("--enable-javascript")
    options.add_experimental_option("excludeSwitches", ['enable-automation'])
    # driver = webdriver.Chrome(service=service, options=options)
    driver = webdriver.Chrome(PATH, options=options)
    return driver

def logging_in(driver):
    url='https://quillbot.com/'
    driver.get(url)
    driver.find_element(By.XPATH,'//*[@id="root-client"]/div[2]/div[2]/div/header/div/div[3]/div/a/button').click()
    user_id = '/html/body/div[2]/div[2]/div[3]/section[1]/div/div/div/div/div/div[1]/div[3]/div[2]/div[2]/div[2]/div[2]/div[2]/div[2]/div/input'
    time.sleep(3)
    driver.find_element(By.XPATH,user_id).click()
    driver.find_element(By.XPATH,user_id).send_keys('rohini.mishra@collegedunia.com')
    time.sleep(2)
    passw= '/html/body/div[2]/div[2]/div[3]/section[1]/div/div/div/div/div/div[1]/div[3]/div[2]/div[2]/div[2]/div[2]/div[3]/div[2]/div/input'
    login = '/html/body/div[2]/div[2]/div[3]/section[1]/div/div/div/div/div/div[1]/div[3]/div[2]/div[2]/div[2]/div[2]/div[5]/button'
    driver.find_element(By.XPATH,passw).click()
    driver.find_element(By.XPATH,passw).send_keys('Rm@123456')
    time.sleep(2)
    driver.find_element(By.XPATH,login).click()
    time.sleep(3)

def get_paraphrased_content(text,driver):
    text_box = '/html/body/div[2]/div[2]/div[3]/section[1]/div/div/div/div/div/div/div/div[3]/div/div/div[2]/div[2]/div/div[1]/div/div[1]/div[1]'
    WebDriverWait(driver, 10).until(EC.element_to_be_clickable((By.XPATH, text_box))).send_keys(text)
    button = '/html/body/div[2]/div[2]/div[3]/section[1]/div/div/div/div/div/div/div/div[3]/div/div/div[2]/div[2]/div/div[1]/div/div[2]/div/div/div[2]/div/button'
    element = WebDriverWait(driver, 10).until(
            EC.element_to_be_clickable((By.XPATH, button)))
    element.click()
    output = '/html/body/div[2]/div[2]/div[3]/section[1]/div/div/div/div/div/div/div/div[3]/div/div/div[2]/div[2]/div/div[2]/div/div[1]'
    para_content = driver.find_element(By.XPATH,output).text
    return para_content

def logging_out(driver):
    profile = '/html/body/div[2]/div[2]/div[2]/div/header/div/div[3]/div/div/span/button'
    profile_ele = driver.find_element(By.XPATH,profile)
    logout = '/html/body/div[8]/div[3]/div/div[2]'
    a = ActionChains(driver)
    a.move_to_element(profile_ele).perform()
    time.sleep(2)
    driver.find_element(By.XPATH,logout).click()

def step_5():
    gc = gspread.service_account(filename = "/home/binoy/course-data-323207-5b61632e094e.json")
    google_sheet = gc.open('second_answers_final')

    sim_data_worksheet = google_sheet.worksheet('similarity_data')
    sim_data = sim_data_worksheet.get_all_records()
    sim_df = pd.DataFrame.from_dict(sim_data)

    comp_data_worksheet = google_sheet.worksheet('competitor_data')
    all_comp_data = comp_data_worksheet.get_all_records()
    comp_df = pd.DataFrame.from_dict(all_comp_data)
    ques_cols = [col for col in comp_df.columns if col.startswith('ques_html')]
    # opt_1_cols = [col for col in comp_df.columns if col.startswith('opt_1_html')]
    # opt_2_cols = [col for col in comp_df.columns if col.startswith('opt_2_html')]
    # opt_3_cols = [col for col in comp_df.columns if col.startswith('opt_3_html')]
    # opt_4_cols = [col for col in comp_df.columns if col.startswith('opt_4_html')]
    sol_cols = [col for col in comp_df.columns if col.startswith('sol_html')]
    if len(ques_cols) > 1:
        ques_cols.sort()
        comp_df['temp'] = ''
        for i in ques_cols:
            comp_df['temp'] = comp_df['temp'] + comp_df[i]
        comp_df.drop(columns=ques_cols,inplace = True)
        comp_df.rename(columns={'temp': 'ques_html'}, inplace=True)
    
    if len(sol_cols) > 1:
        sol_cols.sort()
        comp_df['temp'] = ''
        for i in sol_cols:
            comp_df['temp'] = comp_df['temp'] + comp_df[i]
        comp_df.drop(columns=sol_cols,inplace = True)
        comp_df.rename(columns={'temp': 'ques_html'}, inplace=True)

    df = sim_df.loc[sim_df['verdict'] == 'yes']

    for i in df.index:
        driver = initiate_driver()
        data = {}
        img_details = {}
        cd_id = df['cd_id'][i]
        comp_url = df['comp_url'][i]
        temp = comp_df.loc[(comp_df['cd_id'] == cd_id) & (comp_df['url'] == comp_url)]
        answer_html = temp['sol_html'].iloc[0]

        if df['domain'][i] == 'vedantu':
            prepared_text,img_details,eqn_placeholders = vedantu.vedantu_html_parser(answer_html)
            uploading_content = get_paraphrased_content(prepared_text,driver)
            for i in eqn_placeholders.keys():
                uploading_content.replace(eqn_placeholders[i],i)
        elif df['domain'][i] == 'toppr':
            prepared_text,img_details = toppr.toppr_html_parser(answer_html)
            if prepared_text == 0:
                uploading_content = 'need_to_review'
            else:
                uploading_content = get_paraphrased_content(prepared_text,driver)
            # prepared_text,img_details,eqn_placeholders = toppr.toppr_html_parser(answer_html)
            # uploading_content = get_paraphrased_content(prepared_text,driver)
            # for i in eqn_placeholders.keys():
            #     uploading_content.replace(eqn_placeholders[i],i)
        elif df['domain'][i] == 'byjus':
            prepared_text,img_details = byjus.byjus_html_parser(answer_html)
            if prepared_text == 0:
                uploading_content = 'need_to_review'
            else:
                uploading_content = get_paraphrased_content(prepared_text,driver)
        elif df['domain'][i] == 'embibe':
            prepared_text,img_details = embibe.embibe_html_parser(answer_html)
            if prepared_text == 0:
                uploading_content = 'need_to_review'
            else:
                uploading_content = get_paraphrased_content(prepared_text,driver)
        else:
            uploading_content = 'NA'

        data['cd_id']  = cd_id
        data['comp_url'] = comp_url
        data['uploading_content'] = uploading_content
        data['image_details'] = img_details

        google_sheet.values_append('paraphrased_content', {'valueInputOption': 'RAW'}, {'values': pd.DataFrame([data]).values.tolist()})
        driver.refresh()
